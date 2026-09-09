"""Charte graphique OPT'HELIOS pour les documents Word générés par
l'application (rapport d'audit, checklist terrain...).

Reprend les codes exacts et les règles de mise en forme documentés dans la
compétence `opthelios-notes-techniques` (couleurs mesurées sur le logo
officiel le 20/08/2026, police, tailles, bordures), pour que tous les
livrables .docx de l'application (générés jusqu'ici avec le thème par
défaut de Word, sans logo ni couleurs de marque) soient visuellement
alignés sur les notes techniques OPT'HELIOS existantes.

Ce module ne dépend que de `python-docx` (déjà utilisé par
`domain/docx_service.py`) : pas de nouvelle dépendance.
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

# --------------------------------------------------------------------------
# Charte graphique — codes exacts (voir compétence opthelios-notes-techniques)
# --------------------------------------------------------------------------
NAVY = "1C1E3D"          # Bleu marine : texte fort, bordures, bandeau page de garde
GOLD = "F9B13A"          # Or/doré : sous-titres H2, filets d'en-tête/pied de page
SKY = "8FC4E4"           # Bleu ciel : bandeaux de titre H1
CREAM = "FAEAB6"         # Crème doré : lignes de tableau alternées, encarts
PALE_YELLOW = "FBED94"   # Jaune pâle : encadré nom du site en page de garde
BODY_TEXT = "3A3A3A"     # Gris foncé, couleur de texte courant par défaut

FONT_NAME = "Calibri"

LOGO_PATH = "assets/opthelios_logo.png"

OPTHELIOS_ADDRESS_LINE = "OPT'HELIOS — 471 rue de Pratelmat, 56390 Grand-Champ — contact@opthelios.fr"

# --------------------------------------------------------------------------
# Code couleur "feu tricolore" (demande utilisateur, sept. 2026) : couleurs
# sémantiques ajoutées EN COMPLÉMENT de la charte (pas de recouvrement avec
# NAVY/GOLD/SKY, qui restent réservées à l'identité visuelle OPT'HELIOS),
# pour que le maître d'ouvrage repère d'un coup d'œil, dans le rapport
# transmis, les points qui nécessitent une action de ceux qui sont acquis.
STATUS_RED = "C0392B"      # critique / non conforme : action prioritaire
STATUS_ORANGE = "E67E22"   # majeure / non présent / non vérifiable : à traiter
STATUS_YELLOW = "F2C94C"   # mineure : à surveiller, non urgent
STATUS_BLUE = "5DADE2"     # information : signalé, pas un écart noté comme tel
STATUS_GREEN = "27AE60"    # conforme : acquis
STATUS_GRAY = "95A5A6"     # sans objet / non renseigné : neutre

CRITICITE_COLORS = {
    "critique": STATUS_RED,
    "majeure": STATUS_ORANGE,
    "mineure": STATUS_YELLOW,
    "information": STATUS_BLUE,
}

VERDICT_COLORS = {
    "conforme": STATUS_GREEN,
    "non_conforme": STATUS_RED,
    "non_present": STATUS_ORANGE,
    "non_verifiable": STATUS_ORANGE,
    "sans_objet": STATUS_GRAY,
}

VERDICT_LABELS = {
    "conforme": "Conforme",
    "non_conforme": "Non conforme",
    "non_present": "Non présent",
    "non_verifiable": "Non vérifiable",
    "sans_objet": "Sans objet",
}

CRITICITE_LABELS = {
    "critique": "Critique",
    "majeure": "Majeure",
    "mineure": "Mineure",
    "information": "Information",
}


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def add_status_badge(paragraph, raw_value: str, kind: str = "verdict") -> None:
    """Ajoute une pastille colorée ("●") suivie du libellé, pour le code
    couleur feu tricolore (verdict conforme/non conforme/... ou niveau de
    criticité critique/majeure/mineure/information). `kind` vaut "verdict"
    ou "criticite" ; toute valeur non reconnue retombe sur un gris neutre
    plutôt que de faire échouer la génération du document."""
    value = (raw_value or "").strip().lower()
    if kind == "criticite":
        color = CRITICITE_COLORS.get(value, STATUS_GRAY)
        label = CRITICITE_LABELS.get(value, raw_value or "")
    else:
        color = VERDICT_COLORS.get(value, STATUS_GRAY)
        label = VERDICT_LABELS.get(value, raw_value or "")

    dot = paragraph.add_run("● ")
    dot.font.name = FONT_NAME
    dot.font.size = Pt(10)
    dot.font.color.rgb = _rgb(color)
    dot.font.bold = True

    text = paragraph.add_run(label)
    text.font.name = FONT_NAME
    text.font.size = Pt(10)
    text.font.color.rgb = _rgb(color)
    text.font.bold = True


def add_colored_badge(paragraph, label: str, color_hex: str) -> None:
    """Comme `add_status_badge`, mais pour un statut calculé arbitraire (ex.
    indicateurs de performance FSAV/Prod/Taux, ratio réel/théorique — voir
    `domain/performance_service.py::IndicateurStatut`) plutôt qu'un verdict
    ou une criticité du catalogue de contrôles."""
    dot = paragraph.add_run("● ")
    dot.font.name = FONT_NAME
    dot.font.size = Pt(10)
    dot.font.color.rgb = _rgb(color_hex)
    dot.font.bold = True

    text = paragraph.add_run(label)
    text.font.name = FONT_NAME
    text.font.size = Pt(10)
    text.font.color.rgb = _rgb(color_hex)
    text.font.bold = True


# --------------------------------------------------------------------------
# Primitives bas niveau (bordures de paragraphe, ombrage, champs de page)
# --------------------------------------------------------------------------

def set_paragraph_shading(paragraph, fill_hex: str) -> None:
    """Ombrage plein du paragraphe (bandeau de titre, encart...)."""
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    p_pr.append(shd)


def set_paragraph_border(paragraph, edge: str = "bottom", color: str = GOLD, sz: int = 6) -> None:
    """Filet simple sur un bord du paragraphe (haut/bas), utilisé pour les
    filets or d'en-tête/pied de page et la bordure basse marine des H1."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    border = OxmlElement(f"w:{edge}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(sz))
    border.set(qn("w:space"), "0")
    border.set(qn("w:color"), color)
    p_bdr.append(border)


def add_page_field(paragraph, field_code: str) -> None:
    """Insère un champ Word (ex. PAGE, NUMPAGES) dans le paragraphe donné,
    pour la pagination "Page X / Y" du pied de page."""
    run = paragraph.add_run()
    run.font.name = FONT_NAME
    run.font.size = Pt(7)

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_right_tab(paragraph, usable_width_twips: int = 9350) -> None:
    paragraph.paragraph_format.tab_stops.add_tab_stop(Twips(usable_width_twips), WD_TAB_ALIGNMENT.RIGHT)


def set_table_borders(table, color: str = NAVY, sz: int = 4) -> None:
    """Bordures fines marine sur l'ensemble d'un tableau (au lieu du gris
    par défaut du style "Table Grid")."""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def shade_cell(cell, fill_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def set_cell_border(cell, edge: str, color: str, sz: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    border = OxmlElement(f"w:{edge}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(sz))
    border.set(qn("w:space"), "0")
    border.set(qn("w:color"), color)
    tc_borders.append(border)


def set_cell_margins(cell, twips: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(twips))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


# --------------------------------------------------------------------------
# Mise en page globale du document
# --------------------------------------------------------------------------

def apply_base_document_style(document: Document) -> None:
    """Police, taille et couleur de texte par défaut + format A4 + marges
    charte (1000 twips ≈ 1,76 cm)."""
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(10)
    normal.font.color.rgb = _rgb(BODY_TEXT)
    r_pr = normal.element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), FONT_NAME)

    for style_name, size, bold in [
        ("Title", 20, True),
        ("Heading 1", 13, True),
        ("Heading 2", 11.5, True),
        ("Heading 3", 10.5, True),
    ]:
        if style_name in styles:
            styles[style_name].font.name = FONT_NAME
            styles[style_name].font.size = Pt(size)
            styles[style_name].font.bold = bold

    section = document.sections[0]
    section.page_width = Twips(11907)   # A4 210 mm
    section.page_height = Twips(16840)  # A4 297 mm
    section.top_margin = Twips(1000)
    section.bottom_margin = Twips(1000)
    section.left_margin = Twips(1000)
    section.right_margin = Twips(1000)
    section.header_distance = Twips(500)
    section.footer_distance = Twips(500)


def style_heading_paragraph(paragraph, level: int) -> None:
    """Applique le bandeau/couleur/police charte à un paragraphe déjà créé
    par `document.add_heading(...)` (Heading 1/2/3)."""
    if level == 1:
        set_paragraph_shading(paragraph, SKY)
        set_paragraph_border(paragraph, "bottom", NAVY, 24)
        paragraph.paragraph_format.space_before = Pt(14)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.35
        for run in paragraph.runs:
            run.font.name = FONT_NAME
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = _rgb(NAVY)
    elif level == 2:
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(4)
        for run in paragraph.runs:
            run.font.name = FONT_NAME
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = _rgb(GOLD)
    else:
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(3)
        for run in paragraph.runs:
            run.font.name = FONT_NAME
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = _rgb(NAVY)


def add_heading(document: Document, text: str, level: int = 1):
    """Remplacement direct de `document.add_heading(text, level)` qui
    applique en plus le style de titre charte OPT'HELIOS."""
    paragraph = document.add_heading(text, level=level)
    style_heading_paragraph(paragraph, level)
    return paragraph


# --------------------------------------------------------------------------
# En-tête / pied de page courants (toutes les pages sauf la page de garde)
# --------------------------------------------------------------------------

def apply_running_header_footer(
    document: Document,
    left_title: str,
    right_title: str = "",
) -> None:
    """En-tête "OPT'HELIOS | [titre du rapport]" à gauche / nom du
    client-site à droite, filet or ; pied de page coordonnées OPT'HELIOS +
    pagination "Page X / Y", filet or. La page de garde (première page) est
    laissée vierge via `different_first_page_header_footer`."""
    section = document.sections[0]
    section.different_first_page_header_footer = True

    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.text = ""
    # CORRECTIF : le style "Header" intégré à Word définit ses propres
    # taquets de tabulation par défaut (un centré vers le milieu de la
    # largeur utile). Un taquet ajouté directement sur le paragraphe
    # s'ajoute à ceux du style au lieu de les remplacer : Word/LibreOffice
    # utilise alors le taquet le plus proche, c'est-à-dire celui du style,
    # et le texte de droite atterrit au milieu de la page au lieu d'être
    # aligné à droite. Repasser au style "Normal" (sans taquet hérité)
    # avant d'ajouter notre propre taquet resout le probleme.
    p.style = document.styles["Normal"]
    add_right_tab(p)
    set_paragraph_border(p, "bottom", GOLD, 6)
    p.paragraph_format.space_after = Pt(3)

    run = p.add_run("OPT'HELIOS")
    run.bold = True
    run.font.name = FONT_NAME
    run.font.size = Pt(8)
    run.font.color.rgb = _rgb(NAVY)

    if left_title:
        run = p.add_run(f"   |   {left_title}")
        run.italic = True
        run.font.name = FONT_NAME
        run.font.size = Pt(8)
        run.font.color.rgb = _rgb(BODY_TEXT)

    p.add_run().add_tab()

    if right_title:
        run = p.add_run(right_title)
        run.bold = True
        run.font.name = FONT_NAME
        run.font.size = Pt(8)
        run.font.color.rgb = _rgb(NAVY)

    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.text = ""
    p.style = document.styles["Normal"]  # voir correctif taquet ci-dessus (en-tête)
    add_right_tab(p)
    set_paragraph_border(p, "top", GOLD, 6)
    p.paragraph_format.space_before = Pt(3)

    run = p.add_run(OPTHELIOS_ADDRESS_LINE)
    run.font.name = FONT_NAME
    run.font.size = Pt(7)
    run.font.color.rgb = _rgb(BODY_TEXT)

    p.add_run().add_tab()

    run = p.add_run("Page ")
    run.font.name = FONT_NAME
    run.font.size = Pt(7)
    add_page_field(p, "PAGE")
    run = p.add_run(" / ")
    run.font.name = FONT_NAME
    run.font.size = Pt(7)
    add_page_field(p, "NUMPAGES")

    # Page de garde : header/footer vides (pas de filet, pas de mention
    # "Page 1/N" sur la couverture, comme sur les notes techniques).
    first_header = section.first_page_header
    first_header.is_linked_to_previous = False
    first_header.paragraphs[0].text = ""
    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    first_footer.paragraphs[0].text = ""


# --------------------------------------------------------------------------
# Tableaux
# --------------------------------------------------------------------------

def style_table(table, header_row: int = 1) -> None:
    """Style de tableau charte : en-tête fond marine / texte blanc gras,
    lignes de données alternées blanc / crème doré, bordures fines marine."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, NAVY, 4)

    for row_idx, row in enumerate(table.rows):
        is_header = row_idx < header_row
        for cell in row.cells:
            set_cell_margins(cell, 90)
            if is_header:
                shade_cell(cell, NAVY)
            else:
                shade_cell(cell, CREAM if (row_idx - header_row) % 2 == 0 else "FFFFFF")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = FONT_NAME
                    run.font.size = Pt(9.5)
                    if is_header:
                        run.font.bold = True
                        run.font.color.rgb = _rgb("FFFFFF")
                    else:
                        run.font.color.rgb = _rgb(BODY_TEXT)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


# --------------------------------------------------------------------------
# Encarts ("callout")
# --------------------------------------------------------------------------

def add_callout(document: Document, title: str, paragraphs: Iterable[str]) -> None:
    """Encart charte : tableau à 1 cellule, fond crème doré, bordure gauche
    marine épaisse, bordures haut/bas/droite or fines, titre gras marine."""
    table = document.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    shade_cell(cell, CREAM)
    set_cell_margins(cell, 160)
    set_cell_border(cell, "left", NAVY, 24)
    for edge in ("top", "bottom", "right"):
        set_cell_border(cell, edge, GOLD, 4)

    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(title)
    run.bold = True
    run.font.name = FONT_NAME
    run.font.size = Pt(10)
    run.font.color.rgb = _rgb(NAVY)

    for text in paragraphs:
        p = cell.add_paragraph()
        run = p.add_run(text)
        run.font.name = FONT_NAME
        run.font.size = Pt(9.5)
        run.font.color.rgb = _rgb(BODY_TEXT)

    document.add_paragraph("")


# --------------------------------------------------------------------------
# Page de garde
# --------------------------------------------------------------------------

def add_cover_band(document: Document, title: str, subtitles: Iterable[str]) -> None:
    """Bandeau marine pleine largeur en tête de page de garde (titre blanc
    + sous-titres bleu ciel), comme sur les notes techniques OPT'HELIOS."""
    table = document.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    shade_cell(cell, NAVY)
    set_cell_margins(cell, 220)
    cell.text = ""

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.name = FONT_NAME
    run.font.size = Pt(19)
    run.font.color.rgb = _rgb("FFFFFF")

    for subtitle in subtitles:
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(subtitle)
        run.bold = True
        run.font.name = FONT_NAME
        run.font.size = Pt(10.5)
        run.font.color.rgb = _rgb(SKY)

    document.add_paragraph("")


def add_cover_site_band(document: Document, lines: list[tuple[str, bool, bool]]) -> None:
    """Bandeau jaune pâle sous le titre, avec le nom du site (comme
    l'encadré "CHAUFFERIE ECS SOLAIRE / [SITE]" des notes techniques).

    `lines` : liste de (texte, gras, italique)."""
    table = document.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    shade_cell(cell, PALE_YELLOW)
    set_cell_margins(cell, 200)
    for edge in ("top", "bottom", "left", "right"):
        set_cell_border(cell, edge, GOLD, 4)
    cell.text = ""

    first = True
    for text, bold, italic in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = FONT_NAME
        run.font.size = Pt(15 if bold and not italic else 9.5)
        run.font.color.rgb = _rgb(GOLD if bold else BODY_TEXT)

    document.add_paragraph("")


def add_cover_footer_block(document: Document, logo_path: str = LOGO_PATH) -> None:
    """Filet or + bloc coordonnées OPT'HELIOS (gauche) / logo (droite), en
    bas de page de garde, reprenant exactement la mise en page des notes
    techniques."""
    rule = document.add_paragraph()
    set_paragraph_border(rule, "top", GOLD, 6)

    table = document.add_table(rows=1, cols=2)
    table.autofit = True
    left_cell, right_cell = table.rows[0].cells

    left_cell.text = ""
    p = left_cell.paragraphs[0]
    run = p.add_run("OPT'HELIOS")
    run.bold = True
    run.font.name = FONT_NAME
    run.font.size = Pt(9)
    run.font.color.rgb = _rgb(NAVY)
    p2 = left_cell.add_paragraph()
    run = p2.add_run("471 rue de Pratelmat, 56390 Grand-Champ")
    run.font.name = FONT_NAME
    run.font.size = Pt(9)
    run.font.color.rgb = _rgb(BODY_TEXT)
    p3 = left_cell.add_paragraph()
    run = p3.add_run("contact@opthelios.fr / 06.45.57.10.42")
    run.font.name = FONT_NAME
    run.font.size = Pt(9)
    run.font.color.rgb = _rgb(BODY_TEXT)

    right_cell.text = ""
    if Path(logo_path).exists():
        p = right_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run()
        run.add_picture(logo_path, width=Inches(1.6))
